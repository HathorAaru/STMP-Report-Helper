from datetime import date
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from .data_loader import load_students
from .word_tables import find_table_containing, set_cell_text_preserve_first_run, shade_cell, clear_cell_shading


ATTAINMENT_MAP = {
    "b": "working toward",
    "em-": "working toward",
    "em": "working toward",
    "ex-": "at",
    "ex": "at",
    "exc-": "above",
    "exc": "above",
}


def academic_year(today=None):
    today = today or date.today()
    if today.month >= 9:
        return f"{today.year} - {today.year + 1}"
    return f"{today.year - 1} - {today.year}"


def percent_text(value):
    if value is None:
        return ""
    percent = float(value) * 100
    if abs(percent - round(percent)) < 0.05:
        return f"{round(percent):.0f}%"
    return f"{percent:.1f}%"


def tracker_band(value):
    if value is None:
        return None
    cleaned = "".join(str(value).split()).casefold()
    return ATTAINMENT_MAP.get(cleaned)


def _replace_year(document, year_text):
    for paragraph in document.paragraphs:
        if "2025" in paragraph.text and "2026" in paragraph.text:
            for run in paragraph.runs:
                if "2025" in run.text or "2026" in run.text:
                    run.text = year_text
            return


def _fill_cover(document, student, class_name, teacher_name):
    table = find_table_containing(document, "Name")
    set_cell_text_preserve_first_run(table.cell(0, 1), student.name)
    set_cell_text_preserve_first_run(table.cell(1, 1), class_name)
    set_cell_text_preserve_first_run(table.cell(2, 1), teacher_name)


def _fill_attendance(document, student):
    table = find_table_containing(document, "Attendance and punctuality")
    set_cell_text_preserve_first_run(table.cell(1, 1), percent_text(student.attendance_actual))
    set_cell_text_preserve_first_run(table.cell(1, 2), "95%")
    set_cell_text_preserve_first_run(table.cell(1, 3), percent_text(student.attendance_authorised))
    set_cell_text_preserve_first_run(table.cell(1, 4), percent_text(student.attendance_unauthorised))
    concern = "Y" if student.attendance_actual is not None and float(student.attendance_actual) < 0.95 else "N"
    set_cell_text_preserve_first_run(table.cell(1, 5), concern)


def _highlight_subject(document, subject, value):
    table = find_table_containing(document, subject)
    target = tracker_band(value)
    option_cells = [table.cell(0, 2), table.cell(0, 3), table.cell(0, 4)]
    for cell in option_cells:
        clear_cell_shading(cell)
    if target is None:
        return
    if target == "above":
        shade_cell(table.cell(0, 2))
    elif target == "at":
        shade_cell(table.cell(0, 3))
    elif target == "working toward":
        shade_cell(table.cell(0, 4))


def generate_report(template_path, output_path, student, class_name, teacher_name, today=None):
    document = Document(str(template_path))
    _replace_year(document, academic_year(today))
    _fill_cover(document, student, class_name, teacher_name)
    _fill_attendance(document, student)
    _highlight_subject(document, "Reading", student.reading)
    _highlight_subject(document, "Writing", student.writing)
    _highlight_subject(document, "Mathematics", student.maths)
    document.save(str(output_path))


def generate_all_reports(template_path, attendance_path, attainment_path, output_dir, class_name, teacher_name):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    students, warnings = load_students(attendance_path, attainment_path)
    report_paths = []
    for student in students:
        filename = f"{student.name} Report.docx"
        output_path = output_dir / filename
        generate_report(template_path, output_path, student, class_name, teacher_name)
        report_paths.append(output_path)
    return report_paths, warnings


def generate_zip(template_path, attendance_path, attainment_path, zip_path, class_name, teacher_name):
    zip_path = Path(zip_path)
    work_dir = zip_path.parent / "generated_reports"
    report_paths, warnings = generate_all_reports(template_path, attendance_path, attainment_path, work_dir, class_name, teacher_name)
    with ZipFile(zip_path, "w") as archive:
        for report_path in report_paths:
            archive.write(report_path, report_path.name)
        if warnings:
            archive.writestr("generation_summary.txt", "Warnings:\n" + "\n".join(f"- {warning}" for warning in warnings))
    return zip_path, warnings
